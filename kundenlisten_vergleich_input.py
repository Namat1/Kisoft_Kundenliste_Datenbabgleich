import io
import re
import csv
from collections import OrderedDict

import pandas as pd
import streamlit as st
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter


APP_VERSION = "2026-05-21 Version 8 - intelligente Kisoft-Erkennung"


st.set_page_config(page_title="Kisoft gegen Kundenliste", layout="wide")


STANDARD_SPALTEN = ["SAP", "CSB", "Kundenname", "Straße", "Postleitzahl", "Ort"]


SPALTEN_KANDIDATEN = {
    "SAP": [
        "sap",
        "sap nr",
        "sap-nr",
        "sap nummer",
        "sap-nr debitoren",
        "sap nr debitoren",
        "sap debitoren",
        "sap-nr. debitoren",
        "kundennummer",
        "kunden nummer",
    ],
    "CSB": [
        "csb",
        "csb kundennummer",
        "csb kunden nummer",
        "csb nummer",
        "csb nr",
        "kd nr",
        "kd-nr",
        "kd.-nr",
        "kundennummer",
        "kunden nr",
        "kunden-nr",
        "kdnr",
    ],
    "Kundenname": [
        "kundenname",
        "kunden name",
        "name",
        "kunde",
    ],
    "Straße": [
        "straße",
        "strasse",
        "str",
        "str.",
        "anschrift",
        "adresse",
    ],
    "Postleitzahl": [
        "postleitzahl",
        "plz",
    ],
    "Ort": [
        "ort",
        "stadt",
    ],
}


def text_wert(wert):
    if pd.isna(wert):
        return ""

    text = str(wert).replace("\u00a0", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def nummer_wert(wert):
    text = text_wert(wert)

    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]

    text = re.sub(r"\s+", "", text)
    return text


def sap_wert(wert):
    """
    SAP-Nummern aus Kisoft können führende Nullen haben.
    Beispiel: 0000213003 wird zu 213003.
    """
    text = nummer_wert(wert)

    if re.fullmatch(r"\d+", text):
        text = text.lstrip("0") or "0"

    return text


def vergleich_text(wert):
    """
    Vergleich ohne Beachtung von Groß- und Kleinschreibung.
    """
    text = text_wert(wert)
    text = text.casefold()
    text = re.sub(r"\s+", " ", text).strip()
    return text


def kopf_normalisieren(wert):
    text = text_wert(wert).lower()
    text = text.replace("\n", " ")
    text = text.replace("ß", "ss")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def kandidaten_normalisieren(kandidaten):
    return {kopf_normalisieren(x) for x in kandidaten}


def sortier_sap(wert):
    text = str(wert)
    return (0, int(text)) if text.isdigit() else (1, text)


def finde_spalte(tabelle, ziel, quelle=""):
    """
    Intelligente Spaltenerkennung.

    Wichtig für Kisoft:
    - Kundennummer ist meistens die SAP-Nummer.
    - CSB Kundennummer ist die CSB-Nummer.

    Wichtig für Kundenliste:
    - SAP-Nr oder SAP ist die SAP-Nummer.
    - Kd.-Nr oder Kundennummer kann die CSB-Nummer sein.
    """
    quelle_normalisiert = quelle.casefold()
    spalten = list(tabelle.columns)
    norm_spalten = [(kopf_normalisieren(spalte), spalte) for spalte in spalten]

    def erste_spalte_wenn(bedingung, ausgeschlossene_spalten=None):
        ausgeschlossene_spalten = set(ausgeschlossene_spalten or [])

        for normalisiert, original in norm_spalten:
            if original in ausgeschlossene_spalten:
                continue

            if bedingung(normalisiert):
                return original

        return None

    if ziel == "SAP":
        treffer = erste_spalte_wenn(lambda name: "sap" in name)
        if treffer:
            return treffer

        treffer = erste_spalte_wenn(
            lambda name: (
                name in {"kundennummer", "kunden nummer", "kundennr"}
                or (
                    ("kunden" in name or "kunde" in name)
                    and "nummer" in name
                    and "csb" not in name
                    and "kd nr" not in name
                )
            )
        )
        if treffer:
            return treffer

        kandidaten = []

        for normalisiert, original in norm_spalten:
            if "csb" in normalisiert:
                continue

            if "nummer" in normalisiert or "nr" in normalisiert:
                serie = tabelle[original].dropna().map(nummer_wert).head(80)

                if serie.empty:
                    continue

                zahlen_anteil = serie.map(lambda x: bool(re.fullmatch(r"\d+", x))).mean()
                mittlere_laenge = serie.map(len).median()

                kandidaten.append((zahlen_anteil, mittlere_laenge, original))

        if kandidaten:
            kandidaten.sort(reverse=True)
            return kandidaten[0][2]

    if ziel == "CSB":
        treffer = erste_spalte_wenn(
            lambda name: (
                "csb" in name
                and "a z" not in name
                and "az" not in name
            )
        )
        if treffer:
            return treffer

        treffer = erste_spalte_wenn(
            lambda name: name in {
                "kd nr",
                "kd nummer",
                "kdnr",
                "kd nr debitoren",
                "kunden nr",
                "kunden nummer",
                "kundennummer",
            }
        )
        if treffer:
            return treffer

        treffer = erste_spalte_wenn(
            lambda name: "kd" in name and ("nr" in name or "nummer" in name)
        )
        if treffer:
            return treffer

        if "kisoft" in quelle_normalisiert:
            treffer = erste_spalte_wenn(
                lambda name: "csb" in name and "kund" in name
            )
            if treffer:
                return treffer

    if ziel == "Kundenname":
        return erste_spalte_wenn(
            lambda name: (
                name in {"kundenname", "kunden name", "name", "kunde"}
                or "kundenname" in name
                or "kunden name" in name
            )
        )

    if ziel == "Straße":
        return erste_spalte_wenn(
            lambda name: (
                name in {"strasse", "str", "strasse hausnummer", "anschrift", "adresse"}
                or "strasse" in name
                or "straße" in name
                or "anschrift" in name
                or "adresse" in name
            )
        )

    if ziel == "Postleitzahl":
        return erste_spalte_wenn(
            lambda name: (
                name in {"postleitzahl", "plz"}
                or "postleitzahl" in name
                or name == "plz"
            )
        )

    if ziel == "Ort":
        return erste_spalte_wenn(
            lambda name: name in {"ort", "stadt"} or name.endswith(" ort")
        )

    return None


def lese_csv(datei):
    rohdaten = datei.getvalue()
    letzter_fehler = None

    for encoding in ["utf-8-sig", "utf-8", "cp1252", "latin1"]:
        try:
            text = rohdaten.decode(encoding)
            probe = text[:8192]

            try:
                dialect = csv.Sniffer().sniff(probe, delimiters=";,|\t,")
                trenner = dialect.delimiter
            except Exception:
                trenner = ";"

            return pd.read_csv(io.StringIO(text), sep=trenner, dtype=str)
        except Exception as fehler:
            letzter_fehler = fehler

    raise ValueError(f"CSV konnte nicht gelesen werden: {letzter_fehler}")


def lese_excel(datei):
    excel = pd.ExcelFile(datei)

    bevorzugte_blaetter = []

    for blatt in excel.sheet_names:
        if "liste" in blatt.lower():
            bevorzugte_blaetter.append(blatt)

    bevorzugte_blaetter.extend(
        [blatt for blatt in excel.sheet_names if blatt not in bevorzugte_blaetter]
    )

    beste_tabelle = None
    bestes_blatt = None
    beste_treffer = -1

    for blatt in bevorzugte_blaetter:
        tabelle = pd.read_excel(excel, sheet_name=blatt, dtype=str)

        treffer = sum(
            1
            for ziel in STANDARD_SPALTEN
            if finde_spalte(tabelle, ziel) is not None
        )

        if treffer > beste_treffer:
            beste_treffer = treffer
            beste_tabelle = tabelle
            bestes_blatt = blatt

    if beste_tabelle is None or beste_treffer < 4:
        raise ValueError("In der Excel-Datei wurden nicht genug passende Spalten gefunden.")

    st.caption(f"Excel-Blatt verwendet: {bestes_blatt}")
    return beste_tabelle


def lese_datei(datei):
    name = datei.name.lower()

    if name.endswith(".csv"):
        return lese_csv(datei)

    if name.endswith((".xlsx", ".xlsm", ".xls")):
        return lese_excel(datei)

    raise ValueError("Bitte CSV oder Excel hochladen.")


def baue_standard_tabelle(tabelle, quelle):
    zuordnung = OrderedDict()
    fehlend = []

    for ziel in STANDARD_SPALTEN:
        spalte = finde_spalte(tabelle, ziel, quelle)

        if spalte is None:
            fehlend.append(ziel)
        else:
            zuordnung[ziel] = spalte

    if fehlend:
        raise ValueError(
            f"{quelle}: Diese Spalten wurden nicht gefunden: {', '.join(fehlend)}"
        )

    ausgabe = pd.DataFrame()
    ausgabe["SAP"] = tabelle[zuordnung["SAP"]].map(sap_wert)
    ausgabe["CSB"] = tabelle[zuordnung["CSB"]].map(nummer_wert)
    ausgabe["Kundenname"] = tabelle[zuordnung["Kundenname"]].map(text_wert)
    ausgabe["Straße"] = tabelle[zuordnung["Straße"]].map(text_wert)
    ausgabe["Postleitzahl"] = tabelle[zuordnung["Postleitzahl"]].map(nummer_wert)
    ausgabe["Ort"] = tabelle[zuordnung["Ort"]].map(text_wert)

    ohne_sap = ausgabe[ausgabe["SAP"] == ""].copy()
    ohne_sap.insert(0, "Quelle", quelle)

    ausgabe = ausgabe[ausgabe["SAP"] != ""].copy()
    ausgabe = ausgabe.reset_index(drop=True)

    return ausgabe, ohne_sap, zuordnung


def vergleichsschluessel(tabelle):
    temp = tabelle.copy()

    temp["_CSB"] = temp["CSB"].map(nummer_wert)
    temp["_Kundenname"] = temp["Kundenname"].map(vergleich_text)
    temp["_Straße"] = temp["Straße"].map(vergleich_text)
    temp["_Postleitzahl"] = temp["Postleitzahl"].map(nummer_wert)
    temp["_Ort"] = temp["Ort"].map(vergleich_text)

    return temp


def mache_vergleich(kisoft, kundenliste):
    kisoft_saps = set(kisoft["SAP"])
    kundenliste_saps = set(kundenliste["SAP"])

    nur_kisoft_saps = sorted(kisoft_saps - kundenliste_saps, key=sortier_sap)
    nur_kundenliste_saps = sorted(kundenliste_saps - kisoft_saps, key=sortier_sap)
    gemeinsame_saps = sorted(kisoft_saps & kundenliste_saps, key=sortier_sap)

    nur_kisoft = kisoft[kisoft["SAP"].isin(nur_kisoft_saps)].sort_values("SAP").copy()
    nur_kundenliste = kundenliste[kundenliste["SAP"].isin(nur_kundenliste_saps)].sort_values("SAP").copy()

    kisoft_eindeutig = kisoft.drop_duplicates("SAP", keep="first").set_index("SAP")
    kundenliste_eindeutig = kundenliste.drop_duplicates("SAP", keep="first").set_index("SAP")

    kisoft_vergleich = vergleichsschluessel(kisoft_eindeutig.reset_index()).set_index("SAP")
    kundenliste_vergleich = vergleichsschluessel(kundenliste_eindeutig.reset_index()).set_index("SAP")

    abweichungen = []
    gleiche = 0
    sap_mit_abweichung = set()

    prueffelder = ["CSB", "Kundenname", "Straße", "Postleitzahl", "Ort"]

    for sap in gemeinsame_saps:
        hat_abweichung = False

        for feld in prueffelder:
            vergleichsspalte = "_" + feld

            kisoft_wert = str(kisoft_vergleich.at[sap, vergleichsspalte])
            kundenliste_wert = str(kundenliste_vergleich.at[sap, vergleichsspalte])

            if kisoft_wert != kundenliste_wert:
                hat_abweichung = True
                sap_mit_abweichung.add(sap)

                abweichungen.append(
                    {
                        "SAP": sap,
                        "Feld": feld,
                        "Kisoft-Wert": kisoft_eindeutig.at[sap, feld],
                        "Kundenliste-Wert": kundenliste_eindeutig.at[sap, feld],
                    }
                )

        if not hat_abweichung:
            gleiche += 1

    abweichungen_tabelle = pd.DataFrame(
        abweichungen,
        columns=["SAP", "Feld", "Kisoft-Wert", "Kundenliste-Wert"],
    )

    duplikate_sap = pd.concat(
        [
            kisoft[kisoft.duplicated("SAP", keep=False)].assign(Quelle="Kisoft"),
            kundenliste[kundenliste.duplicated("SAP", keep=False)].assign(Quelle="Kundenliste"),
        ],
        ignore_index=True,
    )

    if not duplikate_sap.empty:
        duplikate_sap = duplikate_sap[["Quelle"] + STANDARD_SPALTEN].sort_values(
            ["SAP", "Quelle"]
        )

    duplikate_csb = pd.concat(
        [
            kisoft[
                (kisoft["CSB"] != "") & kisoft.duplicated("CSB", keep=False)
            ].assign(Quelle="Kisoft"),
            kundenliste[
                (kundenliste["CSB"] != "") & kundenliste.duplicated("CSB", keep=False)
            ].assign(Quelle="Kundenliste"),
        ],
        ignore_index=True,
    )

    if not duplikate_csb.empty:
        duplikate_csb = duplikate_csb[["Quelle"] + STANDARD_SPALTEN].sort_values(
            ["CSB", "Quelle", "SAP"]
        )

    kennzahlen = {
        "Kisoft Zeilen mit SAP": len(kisoft),
        "Kundenliste Zeilen mit SAP": len(kundenliste),
        "Gemeinsame SAP-Nummern": len(gemeinsame_saps),
        "Komplett gleiche gemeinsame SAP-Nummern": gleiche,
        "SAP-Nummern mit echten Abweichungen": len(sap_mit_abweichung),
        "Abweichende Felder gesamt": len(abweichungen_tabelle),
        "Nur in Kisoft - eindeutige SAP": len(nur_kisoft_saps),
        "Nur in Kisoft - Zeilen": len(nur_kisoft),
        "Nur in Kundenliste - eindeutige SAP": len(nur_kundenliste_saps),
        "Nur in Kundenliste - Zeilen": len(nur_kundenliste),
        "Doppelte SAP-Zeilen": len(duplikate_sap),
        "Doppelte CSB-Zeilen": len(duplikate_csb),
    }

    return (
        kennzahlen,
        nur_kisoft,
        nur_kundenliste,
        abweichungen_tabelle,
        duplikate_sap,
        duplikate_csb,
    )


def formatiere_arbeitsmappe(arbeitsmappe):
    dunkel = "1F2937"
    hell = "F9FAFB"
    gelb = "FFF2CC"
    gruen = "E2F0D9"
    rot = "FCE4D6"
    blau = "D9EAF7"

    kopf_fuellung = PatternFill("solid", fgColor=dunkel)
    kopf_schrift = Font(color="FFFFFF", bold=True)
    duenne_linie = Side(style="thin", color="D1D5DB")
    rahmen = Border(
        left=duenne_linie,
        right=duenne_linie,
        top=duenne_linie,
        bottom=duenne_linie,
    )

    for blatt in arbeitsmappe.worksheets:
        blatt.sheet_view.showGridLines = False

        maximale_zeile = blatt.max_row
        maximale_spalte = blatt.max_column

        if blatt.title == "Übersicht":
            blatt["A1"] = "Kisoft gegen Kundenliste"
            blatt["B1"] = APP_VERSION

            for zelle in ["A1", "B1"]:
                blatt[zelle].fill = kopf_fuellung
                blatt[zelle].font = Font(color="FFFFFF", bold=True, size=13)
                blatt[zelle].alignment = Alignment(horizontal="center")

            kopfzeile = 3
            blatt.freeze_panes = "A4"
        else:
            kopfzeile = 1
            blatt.freeze_panes = "A2"

        for zelle in blatt[kopfzeile]:
            zelle.fill = kopf_fuellung
            zelle.font = kopf_schrift
            zelle.alignment = Alignment(horizontal="center", vertical="center")
            zelle.border = rahmen

        for zeile in blatt.iter_rows(
            min_row=kopfzeile + 1,
            max_row=maximale_zeile,
            max_col=maximale_spalte,
        ):
            for zelle in zeile:
                zelle.border = rahmen
                zelle.alignment = Alignment(vertical="top", wrap_text=True)

                if zelle.row % 2 == 0:
                    zelle.fill = PatternFill("solid", fgColor=hell)

        if blatt.title == "Übersicht":
            for zeilennummer in range(4, maximale_zeile + 1):
                pruefung = str(blatt.cell(row=zeilennummer, column=1).value or "")

                if "Abweich" in pruefung:
                    farbe = gelb
                elif "Nur in" in pruefung:
                    farbe = rot
                elif "Gleich" in pruefung or "Gemeinsame" in pruefung:
                    farbe = gruen
                elif "Zeilen" in pruefung:
                    farbe = blau
                else:
                    farbe = None

                if farbe:
                    blatt.cell(row=zeilennummer, column=1).fill = PatternFill(
                        "solid", fgColor=farbe
                    )
                    blatt.cell(row=zeilennummer, column=2).fill = PatternFill(
                        "solid", fgColor=farbe
                    )

        for spaltennummer in range(1, maximale_spalte + 1):
            spaltenbuchstabe = get_column_letter(spaltennummer)
            maximale_laenge = 0

            for zeilennummer in range(1, min(maximale_zeile, 500) + 1):
                wert = blatt.cell(row=zeilennummer, column=spaltennummer).value

                if wert is not None:
                    maximale_laenge = max(maximale_laenge, len(str(wert)))

            breite = min(max(maximale_laenge + 2, 12), 55)
            blatt.column_dimensions[spaltenbuchstabe].width = breite

        if maximale_zeile >= kopfzeile and maximale_spalte >= 1:
            blatt.auto_filter.ref = blatt.dimensions


def excel_download(
    kennzahlen,
    nur_kisoft,
    nur_kundenliste,
    abweichungen,
    duplikate_sap,
    duplikate_csb,
    ohne_sap,
    kisoft,
    kundenliste,
):
    ausgabe = io.BytesIO()

    uebersicht = pd.DataFrame(
        {
            "Prüfung": list(kennzahlen.keys()),
            "Wert": list(kennzahlen.values()),
        }
    )

    with pd.ExcelWriter(ausgabe, engine="openpyxl") as writer:
        uebersicht.to_excel(
            writer,
            sheet_name="Übersicht",
            index=False,
            startrow=2,
        )

        nur_kisoft.to_excel(writer, sheet_name="Nur in Kisoft", index=False)
        nur_kundenliste.to_excel(writer, sheet_name="Nur in Kundenliste", index=False)
        abweichungen.to_excel(writer, sheet_name="Abweichungen", index=False)
        duplikate_sap.to_excel(writer, sheet_name="Duplikate SAP", index=False)
        duplikate_csb.to_excel(writer, sheet_name="Duplikate CSB", index=False)
        ohne_sap.to_excel(writer, sheet_name="Zeilen ohne SAP", index=False)
        kisoft.to_excel(writer, sheet_name="Kisoft aufbereitet", index=False)
        kundenliste.to_excel(writer, sheet_name="Kundenliste aufbereitet", index=False)

        formatiere_arbeitsmappe(writer.book)

    ausgabe.seek(0)
    return ausgabe.getvalue()


st.title("Kisoft gegen Kundenliste")
st.caption(APP_VERSION)

st.write(
    "Lade die Datei Kisoft und die Datei Kundenliste hoch. "
    "Der Vergleich läuft über die SAP-Nummer. "
    "Groß- und Kleinschreibung wird nicht als Unterschied gewertet. "
    "Kisoft wird intelligent erkannt, auch wenn die Spalten anders heißen."
)

links, rechts = st.columns(2)

with links:
    kisoft_datei = st.file_uploader(
        "Kisoft Datei",
        type=["csv", "xlsx", "xlsm", "xls"],
        key="kisoft",
    )

with rechts:
    kundenliste_datei = st.file_uploader(
        "Kundenliste Datei",
        type=["csv", "xlsx", "xlsm", "xls"],
        key="kundenliste",
    )

if kisoft_datei and kundenliste_datei:
    try:
        kisoft_roh = lese_datei(kisoft_datei)
        kundenliste_roh = lese_datei(kundenliste_datei)

        kisoft, kisoft_ohne_sap, kisoft_mapping = baue_standard_tabelle(
            kisoft_roh,
            "Kisoft",
        )

        kundenliste, kundenliste_ohne_sap, kundenliste_mapping = baue_standard_tabelle(
            kundenliste_roh,
            "Kundenliste",
        )

        (
            kennzahlen,
            nur_kisoft,
            nur_kundenliste,
            abweichungen,
            duplikate_sap,
            duplikate_csb,
        ) = mache_vergleich(kisoft, kundenliste)

        ohne_sap = pd.concat(
            [kisoft_ohne_sap, kundenliste_ohne_sap],
            ignore_index=True,
        )

        st.subheader("Ergebnis")

        metrik_spalten = st.columns(5)

        metrik_spalten[0].metric(
            "Nur in Kisoft",
            kennzahlen["Nur in Kisoft - eindeutige SAP"],
        )

        metrik_spalten[1].metric(
            "Nur in Kundenliste",
            kennzahlen["Nur in Kundenliste - eindeutige SAP"],
        )

        metrik_spalten[2].metric(
            "SAP mit Abweichung",
            kennzahlen["SAP-Nummern mit echten Abweichungen"],
        )

        metrik_spalten[3].metric(
            "Abweichende Felder",
            kennzahlen["Abweichende Felder gesamt"],
        )

        metrik_spalten[4].metric(
            "Gleich",
            kennzahlen["Komplett gleiche gemeinsame SAP-Nummern"],
        )

        st.caption(
            "Erkannte Spalten Kisoft: "
            + ", ".join([f"{ziel} = {spalte}" for ziel, spalte in kisoft_mapping.items()])
        )

        st.caption(
            "Erkannte Spalten Kundenliste: "
            + ", ".join(
                [f"{ziel} = {spalte}" for ziel, spalte in kundenliste_mapping.items()]
            )
        )

        excel = excel_download(
            kennzahlen,
            nur_kisoft,
            nur_kundenliste,
            abweichungen,
            duplikate_sap,
            duplikate_csb,
            ohne_sap,
            kisoft,
            kundenliste,
        )

        st.download_button(
            "Excel-Auswertung herunterladen",
            data=excel,
            file_name="kisoft_kundenliste_unterschiede.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs(
            [
                "Abweichungen",
                "Nur in Kisoft",
                "Nur in Kundenliste",
                "Duplikate SAP",
                "Duplikate CSB",
                "Zeilen ohne SAP",
            ]
        )

        with tab1:
            st.dataframe(abweichungen, use_container_width=True, hide_index=True)

        with tab2:
            st.dataframe(nur_kisoft, use_container_width=True, hide_index=True)

        with tab3:
            st.dataframe(nur_kundenliste, use_container_width=True, hide_index=True)

        with tab4:
            st.dataframe(duplikate_sap, use_container_width=True, hide_index=True)

        with tab5:
            st.dataframe(duplikate_csb, use_container_width=True, hide_index=True)

        with tab6:
            st.dataframe(ohne_sap, use_container_width=True, hide_index=True)

    except Exception as fehler:
        st.error(str(fehler))

else:
    st.info("Bitte beide Dateien hochladen.")
